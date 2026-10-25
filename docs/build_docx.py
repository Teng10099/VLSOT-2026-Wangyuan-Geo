"""按附件1模板生成技术报告 docx（python-docx）。保留模板封面与四个一级标题，删除全部“填写说明”。
用法：python build_docx.py 输出.docx [模板.docx]      模板默认取 比赛附件/附件1；配图来自 docs/figs（先运行 make_figs.py）"""
import copy
import os
import re
import sys

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from report_data import ABLATION, CEILING, FACTS, LEADERBOARD, METHOD, PLATFORM, REJECTED, TEAM, WEIGHTS

HERE = os.path.dirname(os.path.abspath(__file__))
FIG = os.path.join(HERE, "figs")
TEMPLATE = os.path.join(HERE, "..", "..", "比赛附件", "附件1：人工智能创新大赛项目技术报告模板.docx")
TODO = "【待填写】"
# 联系电话、VLP 用户名只写进交给主办方的 docx，不进公开仓库：从仓库外的本地 json 读取，没有时留“待填写”
PRIVATE = os.environ.get("VLSOT_REPORT_PRIVATE",
                         os.path.join(HERE, "..", "..", "报名与申请", "报告封面信息.json"))


def private_info():
    import json
    try:
        info = json.load(open(PRIVATE, encoding="utf-8"))
    except OSError:
        info = {}
    return {"phone": info.get("phone") or TODO, "vlp_user": info.get("vlp_user") or TODO}
LATIN, EAST, EAST_H = "Times New Roman", "宋体", "黑体"
J, C, L = WD_ALIGN_PARAGRAPH.JUSTIFY, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT


def set_font(run, size, bold=False, east=EAST, latin=LATIN, italic=False, sub=False):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    if sub:
        run.font.subscript = True
    rpr = run._r.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.insert(0, rf)
    for k in ("w:ascii", "w:hAnsi", "w:cs"):
        rf.set(qn(k), latin)
    rf.set(qn("w:eastAsia"), east)


COUNT = {"fig": 0, "tab": 0}  # 图表编号全文连续


class Writer:
    """anchor 不为 None 时，所有内容依次插在 anchor 段落之前；否则追加到正文末尾。"""

    def __init__(self, doc, anchor):
        self.doc, self.anchor = doc, anchor

    def _new(self):
        return self.anchor.insert_paragraph_before() if self.anchor is not None else self.doc.add_paragraph()

    def p(self, content, size=12, bold=False, align=J, indent=True, before=0, after=0, line=1.5, east=EAST,
          keep_next=False):
        """content: 字符串，或 [(文本, 加粗)] / [(文本, 加粗, 样式字典)] 片段列表。"""
        para = self._new()
        pf = para.paragraph_format
        pf.alignment = align
        pf.line_spacing = line
        pf.space_before, pf.space_after = Pt(before), Pt(after)
        pf.first_line_indent = Pt(size * 2) if indent else Pt(0)
        pf.keep_with_next = keep_next
        pieces = [(content, bold)] if isinstance(content, str) else content
        for piece in pieces:
            text, b = piece[0], piece[1]
            opt = piece[2] if len(piece) > 2 else {}
            set_font(para.add_run(text), opt.get("size", size), b, east=opt.get("east", east),
                     latin=opt.get("latin", LATIN), italic=opt.get("italic", False), sub=opt.get("sub", False))
        return para

    def h2(self, text):
        return self.p(text, size=14, bold=True, indent=False, before=8, after=4, east=EAST_H, keep_next=True)

    def link(self, label, url, size=12):
        """加粗标签后换行写网址：左对齐、不两端对齐，避免网址在连字符处被拆开或字距被拉开。"""
        para = self._new()
        pf = para.paragraph_format
        pf.alignment = L
        pf.line_spacing = 1.5
        pf.space_before = pf.space_after = Pt(0)
        pf.first_line_indent = Pt(size * 2)
        pf.keep_with_next = True
        set_font(para.add_run(label), size, True)
        run = para.add_run()
        run.add_break()
        run.add_text(url)
        set_font(run, size - 0.5)
        return para

    def caption(self, label, text, keep_next=False):
        return self.p([(label + "  ", True), (text, False)], size=10.5, align=C, indent=False, before=3, after=6,
                      line=1.25, keep_next=keep_next)

    def fig(self, name, width_cm, text):
        para = self._new()
        para.paragraph_format.alignment = C
        para.paragraph_format.line_spacing = 1.0
        para.paragraph_format.space_before = Pt(6)
        para.paragraph_format.keep_with_next = True
        para.add_run().add_picture(os.path.join(FIG, name), width=Cm(width_cm))
        COUNT["fig"] += 1
        self.caption(f"图 {COUNT['fig']}", text)

    def table(self, rows, widths, text, highlight=(), size=10.5, left_cols=(0,)):
        assert abs(sum(widths) - 14.6) < 0.25, (text, sum(widths))  # A4 版心宽 14.65 cm
        COUNT["tab"] += 1
        self.caption(f"表 {COUNT['tab']}", text, keep_next=True)
        tbl = self.doc.add_table(rows=len(rows), cols=len(rows[0]))
        if self.anchor is not None:
            self.anchor._p.addprevious(tbl._tbl)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl.autofit = False
        three_line(tbl)
        for c, w in enumerate(widths):
            tbl.columns[c].width = Cm(w)
        for r, row in enumerate(rows):
            for c, val in enumerate(row):
                cell = tbl.cell(r, c)
                cell.width = Cm(widths[c])
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                para = cell.paragraphs[0]
                pf = para.paragraph_format
                pf.alignment = L if c in left_cols else C
                pf.line_spacing = 1.0
                pf.space_before = pf.space_after = Pt(2)
                pf.first_line_indent = Pt(0)
                set_font(para.add_run(str(val)), size, bold=(r == 0 or r in highlight))
                if r in highlight:
                    shade(cell, "E6F4F1")
                if r == 0:
                    cell_border(cell, "bottom", 6)
        # 表后留一点空
        self.p("", size=6, indent=False, line=1.0)
        return tbl

    def equation(self, segments):
        """segments: [(文本, 是否下标)]，居中排的公式行。"""
        para = self._new()
        pf = para.paragraph_format
        pf.alignment = C
        pf.line_spacing = 1.5
        pf.first_line_indent = Pt(0)
        for text, sub in segments:
            set_font(para.add_run(text), 12, italic=False, east=LATIN, sub=sub)
        return para


def three_line(tbl):
    """三线表：上下粗线，表头下细线，无竖线；单元格左右边距 0.1 cm，避免窄列里的 SR@0.5 被拆行。"""
    tblPr = tbl._tbl.tblPr
    for tag in ("w:tblBorders", "w:tblCellMar"):
        old = tblPr.find(qn(tag))
        if old is not None:
            tblPr.remove(old)
    mar = OxmlElement("w:tblCellMar")
    for edge, v in (("top", 0), ("left", 57), ("bottom", 0), ("right", 57)):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:w"), str(v))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tblPr.append(mar)
    b = OxmlElement("w:tblBorders")
    for edge, sz in (("top", 12), ("bottom", 12), ("left", 0), ("right", 0), ("insideH", 0), ("insideV", 0)):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single" if sz else "nil")
        if sz:
            el.set(qn("w:sz"), str(sz))
            el.set(qn("w:space"), "0")
            el.set(qn("w:color"), "000000")
        b.append(el)
    tblPr.append(b)


def cell_border(cell, edge, sz):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = tcPr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tcPr.append(borders)
    el = OxmlElement(f"w:{edge}")
    el.set(qn("w:val"), "single")
    el.set(qn("w:sz"), str(sz))
    el.set(qn("w:space"), "0")
    el.set(qn("w:color"), "000000")
    borders.append(el)


def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def fill_field(p, label, value, underline=True, condense_pt=0.0, newline=False, value_pt=None):
    """封面字段：沿用模板第一个 run 的字体字号，标签 + 值（值加下划线，像填好的横线）。
    condense_pt>0 时值的字距收紧该磅数（成员名单刚好超出一行时用）；
    newline=True 时值另起一行（网址太长、不能在连字符处断开时用），value_pt 为值的字号。"""
    proto = copy.deepcopy(p.runs[0]._r.rPr) if p.runs and p.runs[0]._r.rPr is not None else None
    for r in list(p.runs):
        r._r.getparent().remove(r._r)
    for k, (text, ul, cond) in enumerate(((label, False, 0.0), (value, underline, condense_pt))):
        run = p.add_run()
        if proto is not None:
            if run._r.rPr is not None:
                run._r.remove(run._r.rPr)
            run._r.insert(0, copy.deepcopy(proto))
        if k == 1 and newline:
            run.add_break()
        run.add_text(text)
        run.font.underline = ul
        if k == 1 and value_pt:
            run.font.size = Pt(value_pt)
        if cond > 0:
            sp = OxmlElement("w:spacing")
            sp.set(qn("w:val"), str(-int(round(cond * 20))))
            run._r.get_or_add_rPr().append(sp)


def n(x, d=2):
    return f"{x:.{d}f}"


def rich(text):
    """把 “R_k” 这类写法拆成正文 + 下标片段，供 Writer.p 使用。"""
    pieces, pos = [], 0
    for m in re.finditer(r"([A-Za-zσ])_([a-z])", text):
        pieces.append((text[pos:m.end(1)], False))
        pieces.append((m.group(2), False, {"sub": True}))
        pos = m.end()
    pieces.append((text[pos:], False))
    return [p for p in pieces if p[0]]


def main():
    out_path = sys.argv[1]
    doc = Document(sys.argv[2] if len(sys.argv) > 2 else TEMPLATE)
    P = list(doc.paragraphs)
    priv = private_info()
    print("private fields filled:", {k: v != TODO for k, v in priv.items()})
    heads = {k: P[i] for k, i in (("一", 18), ("二", 27), ("三", 38), ("四", 43))}
    for k, i in (("一", 18), ("二", 27), ("三", 38), ("四", 43)):
        assert P[i].text.replace("\u200b", "").strip().startswith(k + "、"), (i, P[i].text)

    # ---------------- 封面
    for i, label, value, ul in ((7, "团队名称：", TEAM["name"], True),
                                (8, "参赛赛道：", "专项命题挑战赛道  赛道A", False),
                                (9, "项目负责人：", TEAM["leader"], True),
                                (10, "所属学院：", TEAM["school"], True),
                                (11, "联系电话：", priv["phone"], True),
                                (12, "团队成员：", TEAM["members"], True),
                                (13, "开源代码地址：", METHOD["repo"], True),
                                (14, "指导教师：", TEAM["advisor"], True),
                                (15, "填写日期：", TEAM["date"], True)):
        fill_field(P[i], label, value, ul, condense_pt=0.5 if i == 12 else 0.0,
                   newline=(i == 13), value_pt=13 if i == 13 else None)
    # 模板“四、结论”里有一个零宽空格，WPS 渲染成空隙
    for r in P[43].runs:
        r.text = r.text.replace("\u200b", "")

    # ---------------- 删除填写说明与占位空段（不碰含分节符的段落）
    drop = [16, 17, 19] + list(range(20, 27)) + [28] + list(range(29, 38)) + [39, 40, 41, 42, 44] + list(range(45, len(P)))
    for i in drop:
        el = P[i]._element
        xml = el.xml
        if "w:sectPr" in xml or "w:drawing" in xml:
            print("keep paragraph", i, "(sectPr/drawing)")
            continue
        if 'w:type="page"' in xml:
            print("note: paragraph", i, "had a page break")
        el.getparent().remove(el)
    # 一级标题沿用模板的字体字号，只加段前段后距与“与下段同页”
    heads["一"].paragraph_format.page_break_before = True
    for h in heads.values():
        h.paragraph_format.keep_with_next = True
        h.paragraph_format.space_before = Pt(12)
        h.paragraph_format.space_after = Pt(6)

    # ---------------- 排行榜统计（SR@0.5>0 的方法）
    names = ["SR@0.5", "SR@0.7", "AOR", "PR@1.0", "ACE", "PR@0.5"]
    scored = [r for r in LEADERBOARD if r[2] > 0]
    ours = next(r for r in LEADERBOARD if r[1] == METHOD["id"])
    rank, mean = [], []
    for k in range(6):
        vals = [r[2 + k] for r in scored]
        better = sum((v < ours[2 + k]) if names[k] == "ACE" else (v > ours[2 + k]) for v in vals)
        rank.append(better + 1)
        mean.append(sum(vals) / len(vals))
    rank14 = sum(r[2] > ours[2] for r in LEADERBOARD) + 1
    ace_sorted = sorted(r[6] for r in scored)
    second_ace = next(r for r in scored if r[6] == ace_sorted[1])
    ace_gain = 100 * (1 - ours[6] / second_ace[6])
    markov = 100 * (1 - ours[6] / 1.0)

    # ================================================================ 一、榜单成绩
    w = Writer(doc, heads["二"])
    w.p([("VLP 平台用户名：", True), (priv["vlp_user"], False), ("；方法名称：", True),
         (f"{METHOD['short']}（{METHOD['long']}），方法页 {METHOD['url']}。以下成绩为 2026 年 10 月 24 日 16:30"
          "（平台显示时间）提交的 v12f 版本。", False)])
    w.fig("fig_leaderboard.png", 14.6, "VLP 平台 VLSOT 赛道排行榜（2026 年 10 月 25 日截图，按 SR@0.5 排序，红框为本队）")
    rows = [["", *names[:4], "ACE/m", "PR@0.5"],
            ["本方法（v12f）", *[n(ours[2 + k], 4 if names[k] == "ACE" else 2) for k in range(6)]],
            [f"排名（SR@0.5>0 的 {len(scored)} 个方法）", *[str(x) for x in rank]],
            [f"这 {len(scored)} 个方法的均值", *[n(mean[k], 4 if names[k] == "ACE" else 2) for k in range(6)]]]
    w.table(rows, [4.1, 1.75, 1.75, 1.75, 1.75, 1.75, 1.75], "本方法最佳提交的各项指标与排名", highlight=(1,))
    analysis = (f"赛道名称为专项命题挑战赛道A（VLSOT），当前排名{rank14}/{len(LEADERBOARD)}（按SR@0.5）。"
                f"核心优势指标为ACE（平均中心误差），得分{ours[6]:.4f} m，在SR@0.5大于0的{len(scored)}个方法中最低，"
                f"比次低的{second_ace[0]}（{second_ace[6]:.4f} m）低{ace_gain:.0f}%；SR@0.5为{ours[2]:.2f}，高于均值{mean[0]:.2f}。"
                f"SR@0.7（{ours[3]:.2f}，第{rank[1]}）和AOR（{ours[4]:.2f}，第{rank[2]}）是短板，高IoU阈值对中心和尺寸的小误差更敏感。"
                f"PR@1.0/PR@0.5的阈值为0.1 m/0.05 m（若为1 m，由ACE可推出PR@1.0不低于{markov:.0f}%），"
                f"本方法为{ours[5]:.2f}/{ours[7]:.2f}；2025年基线条目的PR可能按1 m/0.5 m计算，不宜直接比较。")
    w.p(analysis)
    n_chars = len(re.sub(r"\s", "", analysis))
    print("section-1 analysis chars:", n_chars)
    assert n_chars <= 300, n_chars  # 模板要求：限 300 字以内

    # ================================================================ 二、技术方案与核心创新点
    w = Writer(doc, heads["三"])
    w.h2("（一）代码与模型权重")
    w.link("参赛代码仓库（GitHub，Apache-2.0 许可）：", METHOD["repo"])
    w.p("README 注明“2026 VLP 挑战赛参赛作品”，附 requirements.txt、运行说明、复现脚本"
        "（scripts/run_final.sh、scripts/run_ablation.sh）和技术报告 PPT。")
    w.link(f"模型权重下载链接（百度网盘，提取码 {WEIGHTS['code']}，分享{WEIGHTS['expiry']}）：", WEIGHTS["url"])
    w.p(f"网盘文件 {WEIGHTS['file']}（{WEIGHTS['size']}），包含 YOLO11x 检测器（COCO 预训练，.pt 与转换后的 .onnx）、"
        "在 TRAIN 上拟合的全部参数（params_v12f.json：打分器权重、地平面系数、框中心偏移、boxfit 边噪声尺度），"
        "以及检测框与颜色概率缓存。")
    w.p([("复现核对：", True),
         ("在仓库代码上重新运行，无论用 TRAIN 重新拟合，还是直接读取参数文件推理（不读取任何 TRAIN 标签），"
          "得到的 10 个结果文件都与提交平台的文件逐字节相同；只用 ONNX 模型推理的核对结果见仓库 README。", False)])

    w.h2("（二）任务与总体思路")
    w.p("VLSOT 任务给定一段 30 帧、1920×1080 的路侧单目视频和一段英文目标描述（json 中另有颜色、运动状态、类别三个字段），"
        "要求逐帧输出目标三维框的 l、w、h、x、y、z（virtuallidar 坐标系）。平台按真实 3D IoU 计算 SR@0.5、SR@0.7 和 AOR，"
        "按中心欧氏误差计算 ACE 和 PR。官方数据包不含相机标定，我们使用源数据集 V2X-Seq-SPD 路侧端公开的标定文件"
        "（官方图像与其原图逐字节一致，按图像名对应），只读取内参与外参，不读取其中任何标注。")
    w.p("我们阅读了数据集的描述生成脚本：描述中的属性先由脚本根据标注计算，再由大语言模型改写成句子，其中有大量度量信息，"
        "包括第 0、14、29 帧目标到相机的水平距离（精确到 0.1 m）、长宽高、航向角、所在九宫格区域、同区域内的名次、截断状态、"
        "最近车辆及其相对方位。单目三维定位最难的是深度，而描述已经给出了距离。因此本方法不训练三维回归网络，"
        "而是把任务拆成两步：先在每帧的 2D 检测框中选出描述所指的目标（目标选择），再把文本给出的距离、尺寸、航向与相机几何结合，"
        "把 2D 框解成 3D 框（三维定位）。两步的误差可以分别统计、分别改进。")

    w.h2("（三）技术框架")
    w.fig("fig_pipeline_v.png", 10.0, "方法流程（按箭头顺序；绿色为两个主要创新环节：候选打分与 boxfit 三维定位）")
    steps = [
        ("（1）文本解析。", "用正则表达式抽取尺寸 l/w/h、三处距离、首段的航向角数值与中、末段的朝向词（left/right/front/back）、"
                        "九宫格区域、名次、截断状态以及最近车辆的相对方位。三处距离过三点做二次插值，得到逐帧距离曲线 r(t)"
                        "（只有两处时线性插值，只有一处时取常数）；航向在三段之间插值。"),
        ("（2）二维检测与颜色。", "YOLO11x（Ultralytics 发布的 COCO 预训练权重，未微调）以 1920 像素输入检测 car、motorcycle、bus、truck 四类，"
                           "置信度不低于 0.05 的框都作为候选。颜色概率由数据集作者标注颜色时所用的 EFN 分类器给出，在关键帧及每隔 3 帧计算。"),
        ("（3）候选打分。", "每个候选框计算 9 个特征：框中心射线与地平面交点的距离和文本距离的相对偏差（4% 以内不计，贴近图像边界的框放宽到 10%）、"
                        "类别是否相符、检测置信度、所在格与文本区域的格距、截断状态是否相符、颜色负对数概率、框的宽高比、同格名次与文本名次之差、"
                        "与最近车辆的相对方位是否相符。打分器为线性模型，以 TRAIN 中与 GT 对应的检测框为正样本、同帧其他框为负样本，"
                        "用 softmax 排序损失训练。"),
        ("（4）全局跟踪。", "在 30 帧上用 Viterbi 动态规划选出一条完整轨迹。一元代价在第 0、14、29 帧使用全部特征，其余帧只用距离、类别和颜色；"
                        "二元代价为 2×(1−相邻帧 2D IoU)，运动目标再加方向代价：相邻帧三维位移与文本航向的夹角越大代价越高（位移小于 0.1 m 时不计）。"),
        ("（5）boxfit 三维定位。", "已知文本尺寸 (l, w, h)、航向 ψ 和该帧的距离 r（由文本距离插值），把中心写成 (x, y) = r(cosθ, sinθ)，只剩方位角 θ 和高度 z 两个未知量。"
                              "对每组 (θ, z)，把三维框的 8 个角点投影到图像并取外接矩形，与检测框比较未被截断的边（离图像边界 10 px 以内的边视为截断，不参与），"
                              "最小化"),
    ]
    for head, body in steps:
        w.p([(head, True), (body, False)])
    w.equation([("E(θ, z) = Σ", False), ("k", True), (" [ (R", False), ("k", True), ("(θ, z) − B", False), ("k", True),
                (") / (σ", False), ("k", True), ("S", False), ("k", True), (" + 1) ]² + [ (z − z", False), ("g", True),
                ("(θ) − h/2) / 0.6 ]²", False)])
    w.p(rich("其中求和只取未截断的边 k，R_k、B_k 为投影矩形和检测框第 k 条边的像素坐标，S_k 为检测框的宽或高，σ_k 为 TRAIN 上统计的边噪声尺度"
             "（GT 中心投影矩形与对应检测框之差的稳健标准差），z_g 为全局地平面高度，0.6 m 为高度先验的标准差。用三级由粗到细的网格搜索"
             "（每级 41×31 个点，范围逐级缩小 8 倍）求最小值。描述缺少距离时，退回框中心射线与地平面的交点。"))
    w.p([("（6）时序平滑。", True),
         ("停车目标对 30 帧中心取中位数；运动目标保持文本距离曲线不变，只对方位角做 ±3 帧的局部线性回归平滑，避免滑动平均在首尾帧产生偏差；"
          "高度取整段中位数；尺寸 l、w、h 直接使用文本值。", False)])
    w.p([("可学习参数。", True),
         ("打分器权重 9 个、全局地平面系数 3 个、框中心偏移 2 个、boxfit 边噪声尺度 4 个，全部只用 227 段 TRAIN 标签拟合；"
          "其余超参数由 TRAIN 上的交叉验证选择。测试阶段只用图像、描述、三个字段和标定，没有使用任何测试标签。", False)])

    w.h2("（四）核心创新点")
    ab = {k: r for k, r in enumerate(ABLATION)}
    ideas = [
        ("1. 文本度量约束代替单目深度估计。",
         f"距离、尺寸、航向直接来自描述，深度不再依赖图像估计；运动目标只平滑方位角、距离保持文本曲线，消除了 v11c 滑动平均在首尾帧约 "
         f"{FACTS['ma_end_bias_m']:.2f} m 的径向偏差。平台 ACE 从第一版的 {PLATFORM[0][7]:.2f} m 降到 {PLATFORM[-1][7]:.3f} m，"
         f"为排行榜上 SR@0.5>0 的方法中最低。"),
        ("2. 复刻数据生成规则的跨模态匹配。",
         f"名次、区域、最近车方位都按数据集生成脚本的定义计算（例如名次按格子左上角到框底边中点的距离排序，相对方位以 3.7 m 为阈值划分），"
         f"颜色使用作者打颜色标签时的 EFN 分类器，使候选框属性与描述的计算方式一致。交叉验证 SR@0.5 从 {ab[0][1]:.2f} 提高到 {ab[2][1]:.2f}。"),
        ("3. 截断感知的三维框投影拟合（boxfit）。",
         f"近处目标的 2D 框中心不是 3D 中心的投影，被图像边界截断时偏差更大。boxfit 用整框投影代替框中心，并剔除截断边。"
         f"在距离 20 m 以内、选中框经常贴边的 {FACTS['near_trunc_clips']} 个 TRAIN 片段上，SR@0.5 从射线法的 {FACTS['near_trunc_sr'][0]}% "
         f"提高到 {FACTS['near_trunc_sr'][1]}%；推广到全部帧后，交叉验证 SR@0.7 从 {ab[6][2]:.2f} 提高到 {ab[7][2]:.2f}，"
         f"中心误差小于 0.1 m 的帧比例从 {ab[6][5]:.2f}% 提高到 {ab[7][5]:.2f}%。"),
        ("4. 评测口径反推与同口径验证。",
         f"数据集附带的评测脚本把每行的 l、w、h、x、y、z 两两排序后当作框的角点。v9 与 v11c 两次提交只在一个片段上不同，"
         f"ACE 变化 {PLATFORM[2][7] - PLATFORM[3][7]:.3f} m，超过按角点解释时可能的最大变化 0.27 m，由此确定平台以 x、y、z 为框中心"
         f"计算真实 3D IoU 与中心误差。本地交叉验证采用同一口径，并修正了训练标签中 GT 与检测框的对应：按中心匹配时 "
         f"{FACTS['gt_mismatch'][1]} 个抽样帧中有 {FACTS['gt_mismatch'][0]} 帧（{100 * FACTS['gt_mismatch'][0] / FACTS['gt_mismatch'][1]:.0f}%）"
         f"配到相邻车辆，改为按 GT 投影矩形的 IoU 匹配。v12f 相对 v11c 的交叉验证提升为 +{ab[8][1] - ab[2][1]:.2f}"
         f"（bootstrap 95% 区间 +{FACTS['boot_v12f_vs_v11c'][0]} ~ +{FACTS['boot_v12f_vs_v11c'][1]}），平台 SR@0.5 提升 "
         f"+{PLATFORM[4][3] - PLATFORM[3][3]:.2f}，方向一致。"),
    ]
    for head, body in ideas:
        w.p([(head, True), (body, False)])

    w.h2("（五）与基线方法及主流方案的区别")
    w.p("排行榜前两名 Mono3DVLT-MT、Mono3DVG-TR 是端到端网络，由图像和文本直接回归三维框；名称带“+backproj”的方法（JointNLT、TransVG、ReSC 等）"
        "先做 2D 视觉定位，再把结果反投影到三维。本方法的差异在于：（1）深度不由图像估计，而由描述中的距离约束，"
        "这是 ACE 明显低于其他方法的主要原因；（2）目标选择不学习隐式的图文匹配，而是按数据生成规则显式计算区域、名次、相对方位和颜色，"
        "每个特征都可以单独检查；（3）三维定位不用框中心反投影，而是拟合整框投影并显式处理截断；（4）可学习参数只有 18 个，"
        "用 227 段 TRAIN 标注即可拟合，检测器和颜色分类器直接使用现成权重，没有微调。")

    # ================================================================ 三、测试案例与结果分析
    w = Writer(doc, heads["四"])
    w.h2("（一）验证方式与消融实验")
    w.p("测试集只有 10 段，平台不返回逐段成绩，每天的提交次数也有限。我们在 227 段 TRAIN 上按目标（sequence, track）分组做 5 折交叉验证，"
        "同一目标的片段不会同时出现在训练折和验证折中；指标实现与平台一致（真实 3D IoU，中心欧氏误差，PR 阈值 0.1 m / 0.05 m）。"
        "主要改动都用按片段的 bootstrap 估计 95% 置信区间，确认提升后才提交。")
    rows = [["配置", "SR@0.5", "SR@0.7", "AOR", "ACE/m", "PR@1.0", "PR@0.5"]]
    for name, sr5, sr7, aor, ace, pr10, pr5 in ABLATION:
        rows.append([name, n(sr5), n(sr7), n(aor), n(ace, 3), n(pr10), n(pr5)])
    w.table(rows, [5.0, 1.6, 1.6, 1.6, 1.6, 1.6, 1.6],
            "5 折交叉验证消融（227 段 TRAIN，每行在上一行基础上增加一项改动；PR@1.0 / PR@0.5 为中心误差 <0.1 m / <0.05 m 的帧比例）",
            highlight=(len(rows) - 1,), size=10)
    d = lambda i, j, k: ABLATION[j][k] - ABLATION[i][k]  # noqa: E731
    w.p(f"前两项改动属于目标选择：按生成规则复刻属性定义（SR@0.5 +{d(0, 1, 1):.2f}），以及换用 EFN 颜色分类器并只用 150 m 内的框计算名次"
        f"（+{d(1, 2, 1):.2f}）。"
        f"第三项修正训练标签中 GT 与检测框的对应，主要提高 SR@0.7（+{d(2, 3, 2):.2f}）和 PR@1.0（+{d(2, 3, 5):.2f}）。"
        f"其余属于三维定位：方位角局部线性平滑再提高 SR@0.7 +{d(3, 4, 2):.2f}；boxfit 带来最大的提升，只用于截断框时，边距从 4 px 放宽到 10 px 使 SR@0.5 "
        f"+{d(5, 6, 1):.2f}（YOLO 的截断框常停在离图像边界 5~10 px 处）；推广到全部帧后 SR@0.7 +{d(6, 7, 2):.2f}、PR@1.0 +{d(6, 7, 5):.2f}。"
        f"最后的二次距离插值与放宽高度先验再提高 +{d(7, 8, 1):.2f}（bootstrap 95% 区间 +{FACTS['boot_v12f_vs_v12d'][0]} ~ "
        f"+{FACTS['boot_v12f_vs_v12d'][1]}）。从 v9 配置到 v12f，交叉验证 SR@0.5 共提高 {d(0, 8, 1):.2f}，ACE 从 "
        f"{ABLATION[0][4]:.3f} m 降到 {ABLATION[8][4]:.3f} m。")

    w.h2("（二）平台提交记录")
    rows = [["版本", "时间", "主要改动", "SR@0.5", "SR@0.7", "AOR", "PR@1.0", "ACE/m", "PR@0.5"]]
    for t, v, desc, sr5, sr7, aor, pr10, ace, pr5 in PLATFORM:
        rows.append([v, t, desc, n(sr5), n(sr7), n(aor), n(pr10), n(ace, 4), n(pr5)])
    w.table(rows, [1.1, 1.2, 3.72, 1.43, 1.43, 1.43, 1.43, 1.43, 1.43],
            "平台测试集的 5 次提交（方法页 301，2026 年 10 月 24 日，平台显示时间）", highlight=(len(rows) - 1,), size=9.5,
            left_cols=(2,))
    a, b = PLATFORM[3], PLATFORM[4]
    w.p(f"v9 加入颜色、航向一致性和全局地平面后，SR@0.5 从早期版本的 15~18 提高到 52。v11c 改了属性定义与颜色模型，测试集 SR@0.5 不变，"
        f"ACE 从 {PLATFORM[2][7]:.2f} m 降到 {a[7]:.2f} m。v12f 换成 boxfit 定位并修正平滑，SR@0.5 +{b[3] - a[3]:.2f}、SR@0.7 +{b[4] - a[4]:.2f}、"
        f"AOR +{b[5] - a[5]:.2f}、PR@1.0 +{b[6] - a[6]:.2f}，ACE 降到 {b[7]:.3f} m；PR@0.5 从 {a[8]:.2f} 降到 {b[8]:.2f}，"
        f"即 300 帧中误差小于 5 cm 的帧从 {round(3 * a[8])} 帧变为 {round(3 * b[8])} 帧。交叉验证中这一项从 {ABLATION[2][6]:.2f} 提高到 "
        f"{ABLATION[8][6]:.2f}，测试集上只差 4 帧，我们认为属于小样本波动。")

    w.h2("（三）代表性测试样例")
    w.fig("fig_cases.png", 14.6, "测试集代表性样例（除注明外均为第 0 帧，括号内为描述给出的距离）。红：选中的检测框；"
                                  "绿：v12f 预测三维框的投影；青：v11c 射线法（同一选择，只换三维定位）。只含我们的预测，不含测试标签")
    cases = [
        ("003235（截断货车，16.5 m）：", "货车被图像左边界和下边界截断，检测框只包含可见部分。射线法把框中心当作三维中心的投影，青框偏向画面内侧；"
                                    "boxfit 只使用未截断的上边和右边，绿框的上边、右边与检测框贴合，三维中心被推到画面外侧。"),
        ("001483（左下角截断的轿车，14.7 m）：", "同样只有上边和右边可用。青框整体偏上、偏右，绿框贴合可见的两条边。"),
        ("000237（面包车，22.2 m）：", "目标未截断，射线法与 boxfit 的结果接近，说明 boxfit 在普通目标上与射线法一致，提升集中在截断和近处目标上。"),
        ("006543（白色轿车，21.7 m）：", "描述给出的高度只有 0.6 m，框很扁。框越扁，同样的高度误差对 IoU 的影响越大，这类目标在 SR@0.7 上容易失分。"),
        ("033447（横穿运动的黑色轿车，42.0 m → 46.8 m）：", "目标从左向右运动。v11c 的滑动平均把首尾帧的位置拖向轨迹中部：第 0 帧青框偏右，第 29 帧青框偏左；"
                                                    "v12f 保持文本距离、只平滑方位角，两帧的绿框都与检测框对齐。"),
    ]
    for head, body in cases:
        w.p([(head, True), (body, False)])

    w.h2("（四）误差上限与局限")
    w.p(f"标注本身限制了 PR 和 SR@0.7 的上限（只用 TRAIN 标签统计）。停车目标的 GT 中心逐帧抖动均值为 {CEILING['parked_jitter_m'] * 100:.1f} cm，"
        f"即使预测恒定在最佳位置，中心误差小于 0.1 m 的帧也只有约 {CEILING['parked_frac_lt10cm'] * 100:.0f}%，小于 0.05 m 的只有约 "
        f"{CEILING['parked_frac_lt5cm'] * 100:.0f}%。标注尺寸也逐帧变化：即使中心完全正确、尺寸用文本值，SR@0.7 的上限约为 "
        f"{CEILING['parked_sr07_cap'] * 100:.1f}%（停车）和 {CEILING['moving_sr07_cap'] * 100:.1f}%（运动），AOR 上限约为 "
        f"{CEILING['parked_aor_cap'] * 100:.1f} 和 {CEILING['moving_aor_cap'] * 100:.1f}。运动目标中间帧的距离只能由三处文本距离插值，"
        f"约 {CEILING['moving_radial_frac_lt10cm'] * 100:.0f}% 的帧径向误差超过 0.1 m。")
    w.p(f"交叉验证 {ABLATION[-1][1]:.2f} 与平台 {PLATFORM[-1][3]:.2f} 相差较大。测试集 10 段中，001483 和 003235 是 15 m 左右、位于左下角的截断目标，"
        f"024386 是 61.5 m 处起始截断的运动目标，047771 和 051657 是约 100 m 远的运动目标（这些信息都来自测试描述）；"
        f"而距离 20 m 以内的截断片段在 TRAIN 中只有 {FACTS['near_trunc_clips']}/227 段。我们推测差距主要来自这些训练集中少见的情形，"
        f"但平台不返回逐段成绩，无法逐段核实。")
    rows = [["改动", "交叉验证结果（相对 v12f 88.41）"]] + [[a_, b_] for a_, b_ in REJECTED]
    w.table(rows, [8.2, 6.4], "v12f 之后试过但未采用的改动", size=10, left_cols=(0, 1))
    w.p("这些改动在交叉验证上都没有可信的提升，因此都没有提交。")

    # ================================================================ 四、结论
    w = Writer(doc, None)
    w.p([("感想与收获。", True),
         (f"第一，先读懂数据再建模。描述生成脚本与评测口径的细节（名次的定义、颜色分类器、IoU 的计算方式）对成绩影响很大："
          f"只复刻属性定义和换用作者的颜色分类器两项，交叉验证 SR@0.5 就提高了 {ABLATION[2][1] - ABLATION[0][1]:.2f}；"
          f"平台上 5 次提交的 SR@0.5 从 {PLATFORM[0][3]:.0f} 提高到 {PLATFORM[-1][3]:.2f}。第二，本地同口径验证比反复提交可靠。"
          "测试集只有 10 段、每天提交次数有限，我们只在交叉验证（附 bootstrap 区间）确认提升后才提交。第三，用上限分析决定何时停止。"
          "统计标注抖动后可知 PR 指标的提升空间有限，v12f 之后的 8 项尝试都没有可信提升，因此没有再提交。", False)])
    w.p([("项目评价。", True),
         (f"本方法 SR@0.5 排名 {rank14}/{len(LEADERBOARD)}，ACE 为 SR@0.5>0 的方法中最低，完成了从文本解析、检测、目标选择、三维定位到生成提交文件的"
          "完整流程，可由检测、颜色缓存和 run_final.sh 三步复现。不足在于 SR@0.7 和 PR 偏低，我们推测近处截断目标和约 100 m 远的小目标"
          "仍是主要失分点；方法依赖描述中的度量信息，描述缺少距离时只能退回地平面射线交点。", False)])
    w.p([("改进建议。", True),
         ("（1）检测框精修：boxfit 对框边位置敏感，可用测试时增强（翻转、多尺度）或实例分割细化框边；"
          "（2）远处目标：对 100 m 外的目标做高分辨率裁剪后再检测，或做亚像素框边估计；"
          "（3）整段联合优化：把 30 帧的 boxfit 代价与运动模型放在同一个优化问题中，代替逐帧拟合后再平滑；"
          "（4）尺寸修正：标注尺寸逐帧变化限制了 SR@0.7，可在更多数据上学习从文本尺寸到标注尺寸的修正。", False)])

    doc.save(out_path)
    print("saved", out_path, "tables", len(doc.tables))


if __name__ == "__main__":
    main()
