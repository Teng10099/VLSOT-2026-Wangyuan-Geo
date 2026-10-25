"""生成 5 页技术报告 PPT（python-pptx，可编辑的原生文本框、表格、图表）。
用法：python build_ppt.py out.pptx      配图来自 docs/figs（先运行 make_figs.py）"""
import os
import sys

from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

from report_data import ABLATION, CEILING, METHOD, PLATFORM, TEAM, WEIGHTS

HERE = os.path.dirname(os.path.abspath(__file__))
FIG = os.path.join(HERE, "figs")
FONT = "Microsoft YaHei"
INK, MUTED, BLUE, TEAL, WARN = (RGBColor(0x10, 0x2A, 0x43), RGBColor(0x52, 0x68, 0x7D), RGBColor(0x13, 0x65, 0xA6),
                                RGBColor(0x0F, 0x76, 0x6E), RGBColor(0x9A, 0x5B, 0x19))
CARD, CARD2, LINE = RGBColor(0xE9, 0xF1, 0xF7), RGBColor(0xE6, 0xF4, 0xF1), RGBColor(0xC9, 0xD6, 0xE2)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]
TOTAL = 5


def style_run(run, size, bold=False, color=INK):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    rpr = run._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea", "a:cs"):
        el = rpr.find(qn(tag))
        if el is None:
            el = etree.SubElement(rpr, qn(tag))
        el.set("typeface", FONT)


def text(slide, x, y, w, h, paras, anchor=MSO_ANCHOR.TOP, align=PP_ALIGN.LEFT, spacing=1.12):
    """paras: [(text 或 [(片段, 加粗, 颜色)], 字号, 加粗, 颜色), ...]"""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.04)
    tf.margin_top = tf.margin_bottom = Inches(0.02)
    for k, (content, size, bold, color) in enumerate(paras):
        p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        p.space_after = Pt(size * 0.35)
        pieces = content if isinstance(content, list) else [(content, bold, color)]
        for piece, b, c in pieces:
            style_run(p.add_run(), size, b, c)
            p.runs[-1].text = piece
    return tb


def box(slide, x, y, w, h, fill=CARD, edge=LINE, radius=0.08, width=1.25):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.adjustments[0] = radius
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    shp.line.color.rgb = edge
    shp.line.width = Pt(width)
    shp.shadow.inherit = False
    shp.text_frame.text = ""
    return shp


def line(slide, x1, y1, x2, y2, arrow=False, color=MUTED, width=2.0):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color
    c.line.width = Pt(width)
    if arrow:
        ln = c.line._get_or_add_ln()
        tail = etree.SubElement(ln, qn("a:tailEnd"))
        tail.set("type", "triangle")
        tail.set("w", "med")
        tail.set("len", "med")
    return c


def frame(title, number, subtitle=None):
    s = prs.slides.add_slide(BLANK)
    text(s, 0.55, 0.32, 12.2, 0.62, [(title, 28, True, INK)])
    bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.58), Inches(0.98), Inches(0.75), Inches(0.06))
    bar.fill.solid()
    bar.fill.fore_color.rgb = TEAL
    bar.line.fill.background()
    if subtitle:
        text(s, 0.55, 1.1, 12.2, 0.42, [(subtitle, 15, False, MUTED)])
    text(s, 0.55, 7.02, 9.0, 0.3, [(f"望远队 · {METHOD['short']} · 长安大学第二届人工智能创新大赛 赛道A（VLSOT）", 10.5, False, MUTED)])
    text(s, 11.6, 7.02, 1.2, 0.3, [(f"{number} / {TOTAL}", 10.5, False, MUTED)], align=PP_ALIGN.RIGHT)
    return s


def notes(slide, content):
    slide.notes_slide.notes_text_frame.text = content


# ================================================================ 1 封面与结论
s = prs.slides.add_slide(BLANK)
text(s, 0.6, 0.45, 7.2, 0.45, [("长安大学第二届人工智能创新大赛 · 专项命题挑战赛 赛道 A（VLSOT）", 15, True, BLUE)])
text(s, 0.6, 1.05, 7.0, 1.6, [("文本约束的几何", 40, True, INK), ("三维视觉语言跟踪", 40, True, INK)], spacing=1.0)
text(s, 0.6, 2.85, 7.0, 0.95, [(f"{METHOD['short']}：把描述里的距离、尺寸和航向代入相机几何，", 16.5, False, MUTED),
                                ("用截断感知的 3D 框投影拟合给出每帧中心", 16.5, False, MUTED)])
cards = [("平台测试 SR@0.5", "61.67", "排行榜第 3 / 14"), ("平台测试 ACE", "0.226 m", "非零成绩方法中最低"),
         ("5 折交叉验证 SR@0.5", "88.41", "227 段训练片段")]
for k, (head, val, foot) in enumerate(cards):
    x = 0.6 + k * 2.42
    box(s, x, 4.05, 2.25, 1.55, fill=CARD2 if k < 2 else CARD, edge=TEAL if k < 2 else LINE)
    text(s, x + 0.12, 4.12, 2.05, 0.35, [(head, 12.5, False, MUTED)])
    text(s, x + 0.12, 4.45, 2.05, 0.6, [(val, 28, True, TEAL if k < 2 else BLUE)])
    text(s, x + 0.12, 5.1, 2.05, 0.4, [(foot, 11.5, False, INK)])
text(s, 0.6, 5.95, 7.4, 0.85, [([("望远队", True, BLUE), (f"　{TEAM['school']}　指导老师：{TEAM['advisor']}", False, INK)], 14, False, INK),
                                (f"队长：{TEAM['leader']}　成员：{TEAM['members']}", 14, False, INK)])
pic = s.shapes.add_picture(os.path.join(FIG, "case_006543_f0_full.png"), Inches(7.75), Inches(1.15), width=Inches(5.05))
s.shapes.add_picture(os.path.join(FIG, "case_003235_f0_zoom.png"), Inches(7.75), Inches(4.12), width=Inches(2.45))
s.shapes.add_picture(os.path.join(FIG, "case_033447_f29_zoom.png"), Inches(10.35), Inches(4.12), width=Inches(2.45))
text(s, 7.75, 6.62, 5.1, 0.4, [("测试集样例。红：选中的检测框　绿：预测 3D 框投影　青：v11c 射线法", 10.5, False, MUTED)])
text(s, 0.6, 7.02, 10.9, 0.3, [([("代码：", True, MUTED), (METHOD["repo"].replace("https://", "") + "（Apache-2.0）", False, MUTED),
                                  ("　权重：", True, MUTED), (f"{WEIGHTS['short']}（提取码 {WEIGHTS['code']}）", False, MUTED)],
                                 10.5, False, MUTED)])
text(s, 11.6, 7.02, 1.2, 0.3, [(f"1 / {TOTAL}", 10.5, False, MUTED)], align=PP_ALIGN.RIGHT)
notes(s, "任务：输入 30 帧单目图像、英文目标描述与相机标定，逐帧输出目标的 l w h x y z（virtuallidar 坐标系）。"
         "数据：TRAIN 227 段 / TEST 10 段，每段 30 帧。"
         f"代码 {METHOD['repo']}；权重 {WEIGHTS['url']}（提取码 {WEIGHTS['code']}）。"
         "官方数据包不含标定，我们使用源数据集 V2X-Seq-SPD 路侧端公开的标定文件，只读内参与外参，不读其中任何标注。"
         "平台成绩取自方法页 http://vlp.chd.edu.cn/tasks/method/301 的 2026-10-24 16:30 提交；排名取自 2026-10-25 抓取的排行榜，"
         "按 SR@0.5 排序第 3（18 个方法中 14 个有成绩）；ACE 在 SR@0.5 大于 0 的 10 个方法中最低。"
         "交叉验证为 227 段 TRAIN 按目标 (sequence, track) 分组的 5 折，指标与平台一致（真实 3D IoU、中心欧氏误差）。")

# ================================================================ 2 方法概述
s = frame("方法概述：文本约束的几何流水线", 2, "不训练三维网络：先选对目标，再用文本给出的度量信息把 2D 框解成 3D 框")
B = [("输入", ["30 帧单目图像 1920×1080", "英文描述 + 颜色/状态/类别字段", "相机标定（V2X-Seq-SPD）"], CARD),
     ("① 文本解析", ["尺寸 l/w/h；第 0/14/29 帧距离", "航向与朝向词、九宫格区域、名次", "截断、最近车及其相对方位"], CARD),
     ("② 二维检测", ["YOLO11x（COCO 预训练）", "置信度 ≥0.05 的车辆框都作候选", "EFN 分类器给出颜色概率"], CARD),
     ("③ 候选打分", ["9 个特征：距离一致、区域、", "名次、颜色、类别、截断、方位…", "线性排序模型（TRAIN 拟合）"], CARD2),
     ("④ 全局跟踪", ["Viterbi 动态规划选整条轨迹", "相邻帧框连续", "运动方向与文本航向一致"], CARD),
     ("⑤ boxfit 三维定位", ["文本尺寸 + 航向 + 距离", "3D 框投影对齐检测框未截断边", "地平面作为高度弱先验"], CARD2),
     ("⑥ 时序平滑", ["停车：30 帧取中位数", "运动：保持文本距离曲线", "方位角做局部线性平滑"], CARD),
     ("输出", ["每帧 l w h x y z", "virtuallidar 坐标系", "result/<视频ID>.txt"], CARD)]
BW, BH, GX = 2.78, 1.62, 0.36
X0, Y1, Y2 = 0.6, 1.72, 4.0
for k, (title, lines, fill) in enumerate(B):
    r, c = divmod(k, 4)
    x, y = X0 + c * (BW + GX), (Y1 if r == 0 else Y2)
    edge = TEAL if fill == CARD2 else BLUE
    box(s, x, y, BW, BH, fill=fill, edge=edge, width=1.5)
    text(s, x + 0.14, y + 0.1, BW - 0.25, 0.42, [(title, 15.5, True, edge)])
    text(s, x + 0.14, y + 0.55, BW - 0.22, 1.05, [(ln, 11.5, False, INK) for ln in lines], spacing=1.0)
    if c < 3:
        line(s, x + BW + 0.03, y + BH / 2, x + BW + GX - 0.03, y + BH / 2, arrow=True)
xe = X0 + 3 * (BW + GX) + BW / 2
xs = X0 + BW / 2
ym = (Y1 + BH + Y2) / 2
line(s, xe, Y1 + BH + 0.02, xe, ym)
line(s, xe, ym, xs, ym)
line(s, xs, ym, xs, Y2 - 0.03, arrow=True)
box(s, 0.6, 5.88, 12.13, 0.98, fill=RGBColor(0xF7, 0xFA, 0xFC), edge=LINE)
text(s, 0.78, 5.95, 11.8, 0.9, [
    ([("学到的参数很少：", True, TEAL), ("打分权重 9 个、地平面 3 个、框中心偏移 2 个、boxfit 边噪声 4 个，只用 TRAIN 标签拟合", False, INK)], 13, False, INK),
    ([("与端到端网络的区别：", True, BLUE), ("距离、尺寸、航向直接代入几何求解，选择和定位两步的误差可以分开度量", False, INK)], 13, False, INK)])
notes(s, "① 文本解析用正则抽取描述中的数值和属性；名次、区域、最近车方位的定义按服务器上数据集生成脚本（01Re/02Re/03Re4color、gpt3.py）复刻。"
         "② YOLO11x 为 Ultralytics 发布的 COCO 预训练权重，未微调；颜色概率来自数据集作者打颜色标签所用的 EFN 分类器。"
         "③ 打分器是 softmax 排序的线性模型，训练样本为 TRAIN 片段中 GT 对应的检测框。④ Viterbi 的转移代价为 2D IoU 与航向一致性。"
         "⑤⑥ 见下一页。源码：vlsot_geo/pipeline3.py、boxfit.py、geo_common.py。")

# ================================================================ 3 核心创新点
s = frame("核心创新点", 3)
IDEAS = [
    ("1  文本度量约束取代单目深度估计", TEAL, [
        "描述给出第 0/14/29 帧距离（0.1 m 精度）与尺寸：中心落在文本距离上，深度不靠图像猜",
        "运动目标保持文本距离曲线、只平滑方位角，消除滑动平均在首尾帧约 0.45 m 的偏差"],
     "平台 ACE 2.53 m → 0.226 m"),
    ("2  复刻数据生成规则的跨模态匹配", BLUE, [
        "按数据集描述生成脚本计算九宫格区域、同格名次（格角到框底边中点距离）、最近车方位（3.7 m 阈值）",
        "颜色用作者打标签时的同款 EFN 分类器；9 个特征 + 线性排序 + Viterbi 全局轨迹"],
     "交叉验证 SR@0.5 77.78 → 81.70"),
    ("3  截断感知的 3D 框投影拟合（boxfit）", TEAL, [
        "已知文本尺寸与航向，求中心使 3D 框投影外接矩形对齐检测框；离图像边界 10 px 内的边不参与",
        "近处截断目标的 2D 框中心不是 3D 中心的投影，沿框中心射线取文本距离会把中心放偏"],
     "近处截断片段 SR@0.5 9.6% → 61.8%"),
    ("4  评测口径反推与同口径验证", BLUE, [
        "v9→v11c 只改 1 段，ACE 降 0.369，超过“角点”口径上限 0.27 → 平台用真实 3D IoU；PR 阈值为 0.1/0.05 m",
        "按目标分组 5 折交叉验证；修正训练 GT 与检测框的匹配（原方法 13% 的帧配到邻车）"],
     "v11c → v12f：CV +6.71，平台 SR@0.5 +9.67"),
]
for k, (title, col, lines, result) in enumerate(IDEAS):
    r, c = divmod(k, 2)
    x, y = 0.6 + c * 6.2, 1.35 + r * 2.78
    box(s, x, y, 5.95, 2.55, fill=CARD2 if col == TEAL else CARD, edge=col, width=1.5)
    text(s, x + 0.2, y + 0.12, 5.6, 0.45, [(title, 17, True, col)])
    text(s, x + 0.2, y + 0.62, 5.6, 1.35, [("· " + ln, 12.5, False, INK) for ln in lines], spacing=1.05)
    text(s, x + 0.2, y + 2.0, 5.6, 0.42, [([("结果：", True, col), (result, True, INK)], 13.5, True, INK)])
notes(s, "数字出处：平台方法页 301（ACE 2.5346 → 0.2259）；交叉验证 scripts/run_ablation.sh（v9 77.78、v11c 81.70）；"
         "截断片段 ana_trunc_cv.py（19 个近处截断片段，射线法 9.6%，boxfit 61.8%）；ana_match.py（3405 个抽样帧中 450 帧中心匹配与投影矩形匹配不一致）。"
         "口径推断：v9 与 v11c 只在 001483 一段不同，两者预测中心相距 5.41 m；若平台把 l w h x y z 当角点，伪中心位移只有 2.70 m，ACE 最多变化 0.27，实际变化 0.369。")

# ================================================================ 4 实验结果与消融
s = frame("实验结果与消融", 4, "左：5 折交叉验证（227 段 TRAIN，按目标分组，真实 3D IoU）；右：平台测试集 5 次提交")
rows = [["配置", "SR@0.5", "SR@0.7", "AOR", "ACE/m", "误差<0.1m"]]
for name, sr5, sr7, aor, ace, pr10, _ in ABLATION:
    rows.append([name, f"{sr5:.2f}", f"{sr7:.2f}", f"{aor:.2f}", f"{ace:.3f}", f"{pr10:.2f}"])
tx, ty, tw, th = 0.6, 1.62, 7.45, 4.55
gf = s.shapes.add_table(len(rows), 6, Inches(tx), Inches(ty), Inches(tw), Inches(th))
tbl = gf.table
widths = [3.35, 0.82, 0.82, 0.82, 0.82, 0.82]
for i, wd in enumerate(widths):
    tbl.columns[i].width = Inches(wd)
for r, row in enumerate(rows):
    tbl.rows[r].height = Inches(th / len(rows))
    for c, val in enumerate(row):
        cell = tbl.cell(r, c)
        cell.text = ""
        cell.margin_left = cell.margin_right = Inches(0.06)
        cell.margin_top = cell.margin_bottom = Inches(0.02)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        cell.fill.solid()
        last = r == len(rows) - 1
        cell.fill.fore_color.rgb = CARD if r == 0 else (CARD2 if last else WHITE)
        p = cell.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.LEFT if c == 0 else PP_ALIGN.CENTER
        run = p.add_run()
        run.text = val
        style_run(run, 11 if c == 0 else 11.5, bold=(r == 0 or last), color=TEAL if (last and c > 0) else INK)
cd = CategoryChartData()
cd.categories = [p[1] for p in PLATFORM]
for name, idx in (("SR@0.5", 3), ("SR@0.7", 4), ("AOR", 5)):
    cd.add_series(name, [round(p[idx], 2) for p in PLATFORM])
ch = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(8.3), Inches(1.55), Inches(4.5), Inches(3.6), cd).chart
ch.has_legend = True
ch.legend.position = XL_LEGEND_POSITION.BOTTOM
ch.legend.include_in_layout = False
ch.font.size = Pt(10.5)
ch.font.name = FONT
plot = ch.plots[0]
plot.gap_width = 60
plot.has_data_labels = True
plot.data_labels.font.size = Pt(8.5)
plot.data_labels.number_format = "0"
plot.data_labels.number_format_is_linked = False
for ser, col in zip(plot.series, (TEAL, BLUE, RGBColor(0x9F, 0xB3, 0xC8))):
    ser.format.fill.solid()
    ser.format.fill.fore_color.rgb = col
ch.value_axis.maximum_scale = 70
ch.value_axis.has_major_gridlines = False
ch.value_axis.tick_labels.font.size = Pt(9)
ch.category_axis.tick_labels.font.size = Pt(10)
ace = " → ".join(f"{p[7]:.2f}" for p in PLATFORM)
v11, v12 = PLATFORM[3], PLATFORM[4]
text(s, 8.3, 5.2, 4.5, 1.0, [
    ([("平台 ACE（m）：", True, TEAL), (ace, False, INK)], 11.5, False, INK),
    ([("v11c → v12f：", True, TEAL), (f"SR@0.5 +{v12[3] - v11[3]:.1f}　SR@0.7 +{v12[4] - v11[4]:.1f}　AOR +{v12[5] - v11[5]:.1f}", False, INK)], 11.5, False, INK),
    (f"PR@1.0 {v11[6]:.2f} → {v12[6]:.2f}　PR@0.5 {v11[8]:.2f} → {v12[8]:.2f}", 11.5, False, INK)], spacing=1.0)
box(s, 0.6, 6.25, 12.13, 0.7, fill=RGBColor(0xFB, 0xF6, 0xEE), edge=RGBColor(0xE4, 0xC9, 0x9E))
text(s, 0.75, 6.3, 11.9, 0.62, [([("试过但未采用（CV SR@0.5，相对 v12f）：", True, WARN),
                                  ("按机位校正系统偏差 −4.07；匀速轨迹拟合 −2.6；射线与 boxfit 加权混合 −0.04 ~ −0.26；"
                                   "尺寸缩放 ≤ +0.2 但 SR@0.7 下降；残差回归 +0.56 但 PR@1.0 −0.7", False, INK)], 11.5, False, INK)],
     spacing=1.0)
notes(s, "左表每行在上一行基础上增加一项改动。第 2、3 行改目标选择（属性定义与颜色），第 4 行修正训练标签里 GT 与检测框的匹配，第 5~9 行改三维定位。"
         "“误差<0.1m”是中心误差小于 0.1 m 的帧比例，对应平台 PR@1.0 列。boxfit 边距从 4 px 改为 10 px 的提升来自 YOLO 截断框常停在离边界 5~10 px 处。"
         "右图来自平台方法页：v1/v2 为早期版本，v9 为 11:07 提交，v11c 为 14:48，v12f 为 16:30。")

# ================================================================ 5 测试样例与误差分析
s = frame("测试样例与误差分析", 5, "红：选中的检测框　绿：v12f 预测 3D 框投影　青：v11c 射线法（同一选择，只换三维定位）")
CASES = [("case_003235_f0_zoom.png", "003235 截断货车：绿框贴合可见的上边和右边"),
         ("case_001483_f0_zoom.png", "001483 左下角截断：射线法框中心偏上偏右"),
         ("case_006543_f0_zoom.png", "006543 h=0.6 的扁框：高度误差直接决定 IoU"),
         ("case_033447_f29_zoom.png", "033447 横穿运动：末帧不再被平滑拖回")]
for k, (fn, cap) in enumerate(CASES):
    x = 0.6 + k * 3.08
    s.shapes.add_picture(os.path.join(FIG, fn), Inches(x), Inches(1.55), width=Inches(2.9), height=Inches(2.9))
    text(s, x, 4.5, 2.95, 0.62, [(cap, 11.5, False, INK)], spacing=1.0)
box(s, 0.6, 5.2, 6.0, 1.7, fill=CARD, edge=LINE)
text(s, 0.78, 5.27, 5.7, 1.6, [
    ("误差上限（只用 TRAIN 标签统计）", 14, True, BLUE),
    # \u00a0：数字与单位之间不换行
    (f"· 停车目标标注逐帧抖动均值 {CEILING['parked_jitter_m'] * 100:.0f}\u00a0cm：中心误差 <10\u00a0cm 的帧最多约 "
     f"{CEILING['parked_frac_lt10cm'] * 100:.0f}%，<5\u00a0cm 最多约 {CEILING['parked_frac_lt5cm'] * 100:.0f}%", 11.5, False, INK),
    (f"· 标注尺寸逐帧变化：中心完全正确时 SR@0.7 上限约 {CEILING['parked_sr07_cap'] * 100:.0f}%（停）/ "
     f"{CEILING['moving_sr07_cap'] * 100:.0f}%（动）", 11.5, False, INK),
    ("· 运动目标中间帧距离只能插值：一半帧径向误差 >0.1 m", 11.5, False, INK)], spacing=1.0)
box(s, 6.73, 5.2, 6.0, 1.7, fill=CARD2, edge=LINE)
text(s, 6.91, 5.27, 5.7, 1.6, [
    ("局限与下一步", 14, True, TEAL),
    ("· 测试集难例集中：两个左下角截断的近处目标（约 15 m）、两个约 100 m 远的运动目标", 11.5, False, INK),
    ("· CV 88.4 与平台 61.7 的差距推测主要来自这些难例（近处截断片段在 TRAIN 中只有 19 / 227 段）；平台不给逐段成绩，无法逐段核实", 11.5, False, INK),
    ("· 计划：检测框测试时增强精修框边、远处目标亚像素定位、整段轨迹联合优化", 11.5, False, INK)], spacing=1.0)
notes(s, "图为测试集片段的检查图（只用测试图像、文本和我们的预测，不含任何测试标签）。"
         "误差上限由 ana_jitter.py 统计（TRAIN 中有三处距离且与 GT 相差 0.3 m 以内的片段）：停车目标看 GT 中心相对片段中位数的偏差；"
         "尺寸上限为预测中心等于 GT、尺寸用文本值时的 IoU 分布。"
         "平台不提供逐段成绩，因此没有针对单个测试片段调参。")

out = sys.argv[1]
prs.save(out)
print("saved", out, "slides", len(prs.slides))
